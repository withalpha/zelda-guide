"""
策略：
1. 用 EasyOCR 读取 cooking/1-5.jpg 中每行的料理名称文字
2. 用图像哈希将 Item_Cook 图标与参考图图标匹配
3. 通过行号关联：图标匹配到第N行 → 第N行的OCR文字就是料理名称
4. 输出 txt + excel
"""

import os, json, cv2, imagehash, openpyxl, easyocr
import numpy as np
from PIL import Image
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE      = r'c:\Users\黄志成\Desktop\VIBE CODING\我的游戏网页'
ITEMS_DIR = os.path.join(BASE, 'Items')
COOK_DIR  = os.path.join(BASE, 'images', 'cooking')

# 网格参数（像素分析得出）
ICON_X0   = 39
ICON_X1   = 109
NAME_X0   = 111
NAME_X1   = 307
ROW_START = 29
ROW_H     = 66

# ── Step 1: OCR 读取所有行的料理名称 ─────────────────────────────────────────
print("=== Step 1: OCR 读取料理名称 ===")
print("  初始化 EasyOCR（支持中文简体）...")
reader = easyocr.Reader(['ch_sim', 'en'], gpu=False, verbose=False)

ref_rows = []   # {img_idx, row, name_text, icon_pil, icon_hash}

for img_idx in range(1, 6):
    path    = os.path.join(COOK_DIR, f'{img_idx}.jpg')
    pil_img = Image.open(path).convert('RGB')
    arr     = np.array(pil_img)
    H       = arr.shape[0]
    
    print(f"\n  处理 {img_idx}.jpg ({H}px 高)...")
    
    row = 0
    y   = ROW_START
    while y + ROW_H <= H:
        # 裁切名称区域
        name_crop = pil_img.crop((NAME_X0, y+2, NAME_X1, y+ROW_H-2))
        name_arr  = np.array(name_crop)
        
        # OCR 识别
        ocr_results = reader.readtext(name_arr, detail=0, paragraph=True)
        name_text   = ' '.join(ocr_results).strip()
        
        # 裁切图标区域
        icon_crop = pil_img.crop((ICON_X0+1, y+1, ICON_X1-1, y+ROW_H-1))
        icon_64   = icon_crop.resize((64, 64), Image.LANCZOS)
        icon_hash = imagehash.phash(icon_64, hash_size=16)
        
        ref_rows.append({
            'img_idx':   img_idx,
            'row':       row,
            'y0':        y,
            'name_text': name_text,
            'icon':      icon_64,
            'hash':      icon_hash,
        })
        
        print(f"    行{row:02d}: [{name_text}]")
        row += 1
        y   += ROW_H

print(f"\n  共读取 {len(ref_rows)} 行")

# ── Step 2: 加载 Item_Cook 图标并计算哈希 ────────────────────────────────────
print("\n=== Step 2: 加载 Item_Cook 图标 ===")

item_files = sorted([
    f for f in os.listdir(ITEMS_DIR)
    if f.startswith('Item_Cook') and f.endswith('.png') and f != 'Item_CookSet.png'
])

items = []
for fname in item_files:
    path    = os.path.join(ITEMS_DIR, fname)
    pil_img = Image.open(path).convert('RGBA')
    
    # 去透明背景
    bg = Image.new('RGB', pil_img.size, (255, 255, 255))
    bg.paste(pil_img, mask=pil_img.split()[3])
    icon_64 = bg.resize((64, 64), Image.LANCZOS)
    
    # 多重哈希
    items.append({
        'filename': fname,
        'img_key':  fname.replace('.png', ''),
        'icon':     icon_64,
        'ph':       imagehash.phash(icon_64, hash_size=16),
        'dh':       imagehash.dhash(icon_64, hash_size=16),
        'wh':       imagehash.whash(icon_64, hash_size=16),
    })

print(f"  共 {len(items)} 个 Item_Cook 图标")

# ── Step 3: 计算参考行的多重哈希 ─────────────────────────────────────────────
print("\n=== Step 3: 计算参考行多重哈希 ===")
for cell in ref_rows:
    cell['ph'] = imagehash.phash(cell['icon'], hash_size=16)
    cell['dh'] = imagehash.dhash(cell['icon'], hash_size=16)
    cell['wh'] = imagehash.whash(cell['icon'], hash_size=16)

# ── Step 4: 用 OpenCV 颜色直方图做精确匹配 ───────────────────────────────────
print("\n=== Step 4: 颜色直方图匹配 ===")

def hsv_hist(pil_img):
    arr = np.array(pil_img.convert('RGB'))
    hsv = cv2.cvtColor(arr, cv2.COLOR_RGB2HSV)
    h   = cv2.calcHist([hsv], [0, 1, 2], None, [18, 16, 16], [0,180, 0,256, 0,256])
    cv2.normalize(h, h)
    return h.flatten()

print("  计算参考行直方图...")
for cell in ref_rows:
    cell['hist'] = hsv_hist(cell['icon'])

print("  计算 Item 直方图...")
for item in items:
    item['hist'] = hsv_hist(item['icon'])

# ── Step 5: 综合匹配 ──────────────────────────────────────────────────────────
print("\n=== Step 5: 综合匹配 ===")

results = []
for item in items:
    best_score = -1
    best_cell  = None
    
    for cell in ref_rows:
        # 颜色直方图相关性（越高越好）
        hist_corr = float(cv2.compareHist(
            item['hist'].reshape(-1,1).astype(np.float32),
            cell['hist'].reshape(-1,1).astype(np.float32),
            cv2.HISTCMP_CORREL
        ))
        
        # 哈希距离（越小越好，转换为相似度）
        ph_sim = 1.0 - (item['ph'] - cell['ph']) / 256.0
        dh_sim = 1.0 - (item['dh'] - cell['dh']) / 256.0
        wh_sim = 1.0 - (item['wh'] - cell['wh']) / 256.0
        
        # 综合得分
        score = hist_corr * 0.5 + ph_sim * 0.2 + dh_sim * 0.15 + wh_sim * 0.15
        
        if score > best_score:
            best_score = score
            best_cell  = cell
    
    results.append({
        'filename':  item['filename'],
        'img_key':   item['img_key'],
        'img_idx':   best_cell['img_idx'],
        'row':       best_cell['row'],
        'name_text': best_cell['name_text'],
        'score':     round(best_score, 4),
    })
    
    conf = '高' if best_score >= 0.6 else ('中' if best_score >= 0.4 else '低')
    print(f"  {item['filename']:<42} → [{best_cell['name_text']}]  得分={best_score:.4f} [{conf}]")

# ── Step 6: 统计 ──────────────────────────────────────────────────────────────
high   = sum(1 for r in results if r['score'] >= 0.6)
medium = sum(1 for r in results if 0.4 <= r['score'] < 0.6)
low    = sum(1 for r in results if r['score'] < 0.4)
print(f"\n匹配统计: 高={high} 中={medium} 低={low} 总={len(results)}")

# ── Step 7: 写 TXT ────────────────────────────────────────────────────────────
txt_path = os.path.join(BASE, 'item_cook_mapping.txt')
with open(txt_path, 'w', encoding='utf-8') as f:
    f.write("塞尔达传说：王国之泪 — Item_Cook 图片与料理名称对照表\n")
    f.write("匹配方法：EasyOCR文字识别 + HSV颜色直方图 + 感知哈希\n")
    f.write("=" * 75 + "\n")
    f.write(f"{'编号':<5} {'文件名':<42} {'料理名称（OCR）':<20} {'参考图':<8} {'得分'}\n")
    f.write("-" * 75 + "\n")
    for i, r in enumerate(results, 1):
        f.write(f"{i:<5} {r['filename']:<42} {r['name_text']:<20} "
                f"{r['img_idx']}.jpg行{r['row']:<3} {r['score']:.4f}\n")
    f.write("-" * 75 + "\n")
    f.write(f"高置信度(≥0.60): {high}  中(0.40-0.59): {medium}  低(<0.40): {low}  总计: {len(results)}\n")

print(f"\nTXT 已保存: {txt_path}")

# ── Step 8: 写 Excel ──────────────────────────────────────────────────────────
xlsx_path = os.path.join(BASE, 'item_cook_mapping.xlsx')
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "料理对照表"

thin = Side(style='thin', color='CCCCCC')
bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr  = Alignment(horizontal='center', vertical='center')
lft  = Alignment(horizontal='left',   vertical='center')

def hdr(cell, text, bg='1F4E79'):
    cell.value     = text
    cell.font      = Font(bold=True, color='FFFFFF', size=10)
    cell.fill      = PatternFill(fill_type='solid', fgColor=bg)
    cell.alignment = ctr
    cell.border    = bdr

for col, w in zip('ABCDEF', [5, 44, 22, 10, 6, 10]):
    ws.column_dimensions[col].width = w

for ci, h_text in enumerate(['编号','Item_Cook文件名','料理名称(OCR识别)','参考图','行','匹配得分'], 1):
    hdr(ws.cell(1, ci), h_text)
ws.row_dimensions[1].height = 22
ws.freeze_panes = 'A2'

fill_hi  = PatternFill(fill_type='solid', fgColor='C6EFCE')
fill_mid = PatternFill(fill_type='solid', fgColor='FFEB9C')
fill_lo  = PatternFill(fill_type='solid', fgColor='FFC7CE')

for i, r in enumerate(results, 1):
    fill = fill_hi if r['score'] >= 0.6 else (fill_mid if r['score'] >= 0.4 else fill_lo)
    ws.append([i, r['filename'], r['name_text'], f"{r['img_idx']}.jpg", r['row'], r['score']])
    rn = ws.max_row
    for ci, al in zip(range(1,7), [ctr, lft, lft, ctr, ctr, ctr]):
        c = ws.cell(rn, ci)
        c.fill = fill; c.border = bdr; c.alignment = al

# Sheet2: 参考图OCR结果（所有行）
ws2 = wb.create_sheet("参考图OCR结果")
for col, w in zip('ABCD', [8, 6, 25, 10]):
    ws2.column_dimensions[col].width = w
for ci, h_text in enumerate(['参考图','行号','OCR识别名称','y坐标'], 1):
    hdr(ws2.cell(1, ci), h_text, '375623')
for cell in ref_rows:
    ws2.append([f"{cell['img_idx']}.jpg", cell['row'], cell['name_text'], cell['y0']])

wb.save(xlsx_path)
print(f"Excel 已保存: {xlsx_path}")
print("\n完成！")
