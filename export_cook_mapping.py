"""
从 correct_dishes_mapping.json 导出 Item_Cook 料理对照表
输出：item_cook_mapping.txt 和 item_cook_mapping.xlsx
"""

import os, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = r'c:\Users\黄志成\Desktop\VIBE CODING\我的游戏网页'

# ── 加载映射 ──────────────────────────────────────────────────────────────────
with open(os.path.join(BASE, 'correct_dishes_mapping.json'), encoding='utf-8') as f:
    all_entries = json.load(f)

# 只取 Item_Cook 开头的（排除 Item_Roast / Item_Chilled 等）
cook_entries = [e for e in all_entries if e['img'].startswith('Item_Cook')]
print(f"Item_Cook 条目共 {len(cook_entries)} 条")

# ── 写 TXT ────────────────────────────────────────────────────────────────────
txt_path = os.path.join(BASE, 'item_cook_mapping.txt')
with open(txt_path, 'w', encoding='utf-8') as f:
    f.write("塞尔达传说：王国之泪 — Item_Cook 图片与料理名称对照表\n")
    f.write("=" * 65 + "\n")
    f.write(f"{'编号':<6} {'图片文件名（不含.png）':<42} {'料理名称'}\n")
    f.write("-" * 65 + "\n")
    for e in cook_entries:
        f.write(f"{e['no']:<6} {e['img']:<42} {e['name']}\n")
    f.write("-" * 65 + "\n")
    f.write(f"共 {len(cook_entries)} 条料理\n")

print(f"TXT 已保存: {txt_path}")

# ── 写 Excel ──────────────────────────────────────────────────────────────────
xlsx_path = os.path.join(BASE, 'item_cook_mapping.xlsx')
wb = openpyxl.Workbook()

# ── 辅助样式 ──────────────────────────────────────────────────────────────────
thin  = Side(style='thin', color='CCCCCC')
bdr   = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr   = Alignment(horizontal='center', vertical='center')
left  = Alignment(horizontal='left',   vertical='center')

def hdr_cell(cell, text, bg='1F4E79'):
    cell.value     = text
    cell.font      = Font(bold=True, color='FFFFFF', size=11)
    cell.fill      = PatternFill(fill_type='solid', fgColor=bg)
    cell.alignment = ctr
    cell.border    = bdr

# ── 按字母分组的颜色（A-S 各一色，循环） ─────────────────────────────────────
GROUP_COLORS = [
    'D9E1F2','FCE4D6','E2EFDA','FFF2CC','DDEBF7',
    'F4CCCC','D9D2E9','CFE2F3','D0E0E3','EAD1DC',
    'FFF3CD','E8F5E9','F3E5F5','E0F2F1','FBE9E7',
    'E8EAF6','F9FBE7','FCF3CF','E3F2FD',
]

def get_group_color(img_key):
    # 取字母部分 A-S
    parts = img_key.split('_')  # ['Item','Cook','A','01']
    if len(parts) >= 3:
        letter = parts[2]
        idx = ord(letter[0]) - ord('A')
        return GROUP_COLORS[idx % len(GROUP_COLORS)]
    return 'FFFFFF'

# ── Sheet 1: 完整对照表 ───────────────────────────────────────────────────────
ws = wb.active
ws.title = "Item_Cook 料理对照表"

# 列宽
for col, width in zip('ABCD', [6, 45, 22, 10]):
    ws.column_dimensions[col].width = width

# 表头
hdr_cell(ws['A1'], '编号')
hdr_cell(ws['B1'], 'Item_Cook 图片文件名（不含 .png）')
hdr_cell(ws['C1'], '料理名称')
hdr_cell(ws['D1'], '分组')
ws.row_dimensions[1].height = 24
ws.freeze_panes = 'A2'

for e in cook_entries:
    img_key = e['img']
    # 分组字母
    parts = img_key.split('_')
    group = parts[2] if len(parts) >= 3 else '?'
    # 特殊处理 C_17_xxx → 分组 C
    
    color = get_group_color(img_key)
    fill  = PatternFill(fill_type='solid', fgColor=color)
    
    ws.append([e['no'], img_key, e['name'], group])
    rn = ws.max_row
    for ci, al in zip(range(1, 5), [ctr, left, left, ctr]):
        cell = ws.cell(row=rn, column=ci)
        cell.fill      = fill
        cell.border    = bdr
        cell.alignment = al

# ── Sheet 2: 按分组汇总 ───────────────────────────────────────────────────────
ws2 = wb.create_sheet("按分组汇总")
for col, width in zip('ABCDE', [6, 45, 22, 10, 12]):
    ws2.column_dimensions[col].width = width

hdr_cell(ws2['A1'], '编号', '375623')
hdr_cell(ws2['B1'], 'Item_Cook 图片文件名', '375623')
hdr_cell(ws2['C1'], '料理名称', '375623')
hdr_cell(ws2['D1'], '分组', '375623')
hdr_cell(ws2['E1'], '分组说明', '375623')
ws2.row_dimensions[1].height = 24
ws2.freeze_panes = 'A2'

# 分组说明（根据游戏内容）
GROUP_DESC = {
    'A': '烤/蒸类',
    'B': '炒/煎类',
    'C': '岩盐/药水类',
    'D': '盐烤/香辣/油炸类',
    'E': '油炸/干煎类',
    'F': '炖水果类',
    'G': '汤/炖菜类',
    'H': '炖鱼/咖喱类',
    'I': '咖喱/炖肉/烩饭类',
    'J': '烩饭/菜饭/炒饭类',
    'K': '炒饭/盖饭/饭团类',
    'L': '饭团/粥/面包类',
    'M': '起司面包类',
    'N': '披萨/派类',
    'O': '焗烤/煎蛋卷类',
    'P': '煎蛋卷/甜点类',
    'Q': '甜点/布丁/可丽饼/蛋糕类',
    'R': '蛋糕/派/特殊料理类',
    'S': '怪物/魔人/特殊料理类',
}

for e in cook_entries:
    img_key = e['img']
    parts   = img_key.split('_')
    group   = parts[2] if len(parts) >= 3 else '?'
    desc    = GROUP_DESC.get(group, '')
    color   = get_group_color(img_key)
    fill    = PatternFill(fill_type='solid', fgColor=color)
    
    ws2.append([e['no'], img_key, e['name'], group, desc])
    rn = ws2.max_row
    for ci, al in zip(range(1, 6), [ctr, left, left, ctr, left]):
        cell = ws2.cell(row=rn, column=ci)
        cell.fill      = fill
        cell.border    = bdr
        cell.alignment = al

# ── Sheet 3: 分组统计 ─────────────────────────────────────────────────────────
ws3 = wb.create_sheet("分组统计")
for col, width in zip('ABCD', [8, 20, 10, 30]):
    ws3.column_dimensions[col].width = width

hdr_cell(ws3['A1'], '分组', '7B2D8B')
hdr_cell(ws3['B1'], '分组说明', '7B2D8B')
hdr_cell(ws3['C1'], '料理数量', '7B2D8B')
hdr_cell(ws3['D1'], '料理名称列表', '7B2D8B')
ws3.row_dimensions[1].height = 24

from collections import defaultdict
groups = defaultdict(list)
for e in cook_entries:
    parts = e['img'].split('_')
    g = parts[2] if len(parts) >= 3 else '?'
    groups[g].append(e['name'])

for g in sorted(groups.keys()):
    names = groups[g]
    color = get_group_color(f'Item_Cook_{g}_01')
    fill  = PatternFill(fill_type='solid', fgColor=color)
    desc  = GROUP_DESC.get(g, '')
    
    ws3.append([g, desc, len(names), '、'.join(names)])
    rn = ws3.max_row
    for ci, al in zip(range(1, 5), [ctr, left, ctr, left]):
        cell = ws3.cell(row=rn, column=ci)
        cell.fill      = fill
        cell.border    = bdr
        cell.alignment = al
    ws3.row_dimensions[rn].height = 18

# 总计行
total_row = ws3.max_row + 1
ws3.cell(total_row, 1, '合计').font = Font(bold=True)
ws3.cell(total_row, 3, len(cook_entries)).font = Font(bold=True)
ws3.cell(total_row, 1).alignment = ctr
ws3.cell(total_row, 3).alignment = ctr

wb.save(xlsx_path)
print(f"Excel 已保存: {xlsx_path}")
print(f"\n完成！共导出 {len(cook_entries)} 条 Item_Cook 料理对照")
