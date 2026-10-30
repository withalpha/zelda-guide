"""
用图像模板匹配将 Item_Cook*.png 与 cooking/1-5.jpg 做精确匹配
输出：item_cook_mapping.txt 和 item_cook_mapping.xlsx
"""

import os, json, cv2, imagehash, openpyxl
import numpy as np
from PIL import Image
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── 路径 ──────────────────────────────────────────────────────────────────────
BASE      = r'c:\Users\黄志成\Desktop\VIBE CODING\我的游戏网页'
ITEMS_DIR = os.path.join(BASE, 'Items')
COOK_DIR  = os.path.join(BASE, 'images', 'cooking')

# ── 加载料理名称映射 ──────────────────────────────────────────────────────────
with open(os.path.join(BASE, 'correct_dishes_mapping.json'), encoding='utf-8') as f:
    raw = json.load(f)

# img_key (不含.png) → name
name_map = {entry['img']: entry['name'] for entry in raw}

# ── 加载参考图（灰度 + 彩色） ─────────────────────────────────────────────────
print("加载参考图...")
ref_color = {}
ref_gray  = {}
ref_sizes = {}
for i in range(1, 6):
    path = os.path.join(COOK_DIR, f'{i}.jpg')
    # 用 PIL 读取（支持中文路径），再转 numpy/OpenCV 格式
    pil_ref = Image.open(path).convert('RGB')
    img_rgb = np.array(pil_ref, dtype=np.uint8)
    img_bgr = cv2.cvtColor(img_rgb, cv2.COLOR_RGB2BGR)
    ref_color[i] = img_bgr
    ref_gray[i]  = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2GRAY)
    ref_sizes[i] = (img_bgr.shape[1], img_bgr.shape[0])  # (w, h)
    print(f"  {i}.jpg: {ref_sizes[i]}")

# ── 获取所有 Item_Cook 文件（排除 Item_CookSet.png） ──────────────────────────
item_files = sorted([
    f for f in os.listdir(ITEMS_DIR)
    if f.startswith('Item_Cook') and f.endswith('.png') and f != 'Item_CookSet.png'
])
print(f"\n共 {len(item_files)} 个 Item_Cook 图片")

# ── 模板匹配函数 ──────────────────────────────────────────────────────────────
def match_template_in_refs(item_path):
    """
    在5张参考图中用多尺度模板匹配，返回最佳匹配结果
    """
    # 读取 Item 图片（RGBA → 去透明背景 → 灰度）
    pil_img = Image.open(item_path).convert('RGBA')
    
    # 合成白色背景
    bg = Image.new('RGB', pil_img.size, (255, 255, 255))
    bg.paste(pil_img, mask=pil_img.split()[3])
    
    # 转 OpenCV 格式
    item_bgr  = cv2.cvtColor(np.array(bg), cv2.COLOR_RGB2BGR)
    item_gray = cv2.cvtColor(item_bgr, cv2.COLOR_BGR2GRAY)
    
    ih, iw = item_gray.shape
    
    best = {'score': -1, 'img_idx': 0, 'x': 0, 'y': 0, 'scale': 1.0, 'match_w': iw, 'match_h': ih}
    
    # 多尺度匹配
    for scale in [0.5, 0.6, 0.7, 0.75, 0.8, 0.85, 0.9, 1.0]:
        new_w = max(8, int(iw * scale))
        new_h = max(8, int(ih * scale))
        tmpl  = cv2.resize(item_gray, (new_w, new_h))
        
        for idx in range(1, 6):
            rg = ref_gray[idx]
            if tmpl.shape[0] > rg.shape[0] or tmpl.shape[1] > rg.shape[1]:
                continue
            
            res = cv2.matchTemplate(rg, tmpl, cv2.TM_CCOEFF_NORMED)
            _, max_val, _, max_loc = cv2.minMaxLoc(res)
            
            if max_val > best['score']:
                best['score']   = max_val
                best['img_idx'] = idx
                best['x']       = max_loc[0]
                best['y']       = max_loc[1]
                best['scale']   = scale
                best['match_w'] = new_w
                best['match_h'] = new_h
    
    return best

# ── 执行匹配 ──────────────────────────────────────────────────────────────────
print("\n开始模板匹配...")
results = []

for i, fname in enumerate(item_files):
    item_path = os.path.join(ITEMS_DIR, fname)
    img_key   = fname.replace('.png', '')
    dish_name = name_map.get(img_key, '（未知）')
    
    best = match_template_in_refs(item_path)
    
    # 根据匹配位置推算在参考图中的行列（每格约150x132）
    col = best['x'] // 150
    row = best['y'] // 132
    
    results.append({
        'no':       i + 1,
        'filename': fname,
        'img_key':  img_key,
        'name':     dish_name,
        'img_idx':  best['img_idx'],
        'row':      row,
        'col':      col,
        'x':        best['x'],
        'y':        best['y'],
        'score':    round(best['score'], 4),
        'scale':    best['scale'],
    })
    
    status = '✓' if best['score'] >= 0.7 else ('△' if best['score'] >= 0.5 else '✗')
    print(f"  [{i+1:3d}/{len(item_files)}] {status} {fname:<35} → {best['img_idx']}.jpg "
          f"行{row} 列{col}  得分={best['score']:.3f}  {dish_name}")

# ── 统计 ──────────────────────────────────────────────────────────────────────
high   = sum(1 for r in results if r['score'] >= 0.7)
medium = sum(1 for r in results if 0.5 <= r['score'] < 0.7)
low    = sum(1 for r in results if r['score'] < 0.5)

print(f"\n匹配统计:")
print(f"  高置信度 (≥0.70): {high} 个")
print(f"  中置信度 (0.50-0.69): {medium} 个")
print(f"  低置信度 (<0.50): {low} 个")
print(f"  总计: {len(results)} 个")

# ── 写 TXT ────────────────────────────────────────────────────────────────────
txt_path = os.path.join(BASE, 'item_cook_mapping.txt')
with open(txt_path, 'w', encoding='utf-8') as f:
    f.write("塞尔达传说：王国之泪 — Item_Cook 图片与料理名称对照表\n")
    f.write("=" * 80 + "\n")
    f.write(f"{'编号':<5} {'文件名':<38} {'料理名称':<20} {'参考图':<8} {'行':<4} {'列':<4} {'匹配得分'}\n")
    f.write("-" * 80 + "\n")
    
    for r in results:
        conf = '高' if r['score'] >= 0.7 else ('中' if r['score'] >= 0.5 else '低')
        f.write(f"{r['no']:<5} {r['filename']:<38} {r['name']:<20} "
                f"{r['img_idx']}.jpg  {r['row']:<4} {r['col']:<4} {r['score']:.4f} [{conf}]\n")
    
    f.write("\n" + "=" * 80 + "\n")
    f.write(f"高置信度 (≥0.70): {high} 个\n")
    f.write(f"中置信度 (0.50-0.69): {medium} 个\n")
    f.write(f"低置信度 (<0.50): {low} 个\n")
    f.write(f"总计: {len(results)} 个\n")

print(f"\nTXT 已保存: {txt_path}")

# ── 写 Excel ──────────────────────────────────────────────────────────────────
xlsx_path = os.path.join(BASE, 'item_cook_mapping.xlsx')
wb = openpyxl.Workbook()

# ── Sheet 1: 完整对照表 ───────────────────────────────────────────────────────
ws = wb.active
ws.title = "料理对照表"

# 列宽
col_widths = [6, 42, 22, 10, 6, 6, 12, 10]
for ci, cw in enumerate(col_widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = cw

# 表头
headers = ['编号', 'Item_Cook 文件名', '料理名称', '参考图', '行', '列', '匹配得分', '置信度']
ws.append(headers)

hdr_font  = Font(bold=True, color='FFFFFF', size=11)
hdr_fill  = PatternFill(fill_type='solid', fgColor='1F4E79')
hdr_align = Alignment(horizontal='center', vertical='center')
thin      = Side(style='thin', color='AAAAAA')
border    = Border(left=thin, right=thin, top=thin, bottom=thin)

for cell in ws[1]:
    cell.font      = hdr_font
    cell.fill      = hdr_fill
    cell.alignment = hdr_align
    cell.border    = border

ws.row_dimensions[1].height = 22

# 数据行
fill_high   = PatternFill(fill_type='solid', fgColor='C6EFCE')  # 绿
fill_medium = PatternFill(fill_type='solid', fgColor='FFEB9C')  # 黄
fill_low    = PatternFill(fill_type='solid', fgColor='FFC7CE')  # 红

for r in results:
    conf  = '高' if r['score'] >= 0.7 else ('中' if r['score'] >= 0.5 else '低')
    fill  = fill_high if r['score'] >= 0.7 else (fill_medium if r['score'] >= 0.5 else fill_low)
    
    row_data = [
        r['no'],
        r['filename'],
        r['name'],
        f"{r['img_idx']}.jpg",
        r['row'],
        r['col'],
        r['score'],
        conf,
    ]
    ws.append(row_data)
    
    rn = ws.max_row
    for ci in range(1, 9):
        cell = ws.cell(row=rn, column=ci)
        cell.fill      = fill
        cell.border    = border
        cell.alignment = Alignment(horizontal='center' if ci not in (2, 3) else 'left',
                                   vertical='center')

# 冻结首行
ws.freeze_panes = 'A2'

# ── Sheet 2: 按参考图分组 ─────────────────────────────────────────────────────
ws2 = wb.create_sheet("按参考图分组")
ws2.column_dimensions['A'].width = 10
ws2.column_dimensions['B'].width = 42
ws2.column_dimensions['C'].width = 22
ws2.column_dimensions['D'].width = 12

ws2.append(['参考图', 'Item_Cook 文件名', '料理名称', '匹配得分'])
for cell in ws2[1]:
    cell.font      = Font(bold=True, color='FFFFFF')
    cell.fill      = PatternFill(fill_type='solid', fgColor='375623')
    cell.alignment = Alignment(horizontal='center')

# 按参考图编号排序
for img_idx in range(1, 6):
    group = [r for r in results if r['img_idx'] == img_idx]
    group.sort(key=lambda x: (x['row'], x['col']))
    for r in group:
        conf = '高' if r['score'] >= 0.7 else ('中' if r['score'] >= 0.5 else '低')
        fill = fill_high if r['score'] >= 0.7 else (fill_medium if r['score'] >= 0.5 else fill_low)
        ws2.append([f"{img_idx}.jpg", r['filename'], r['name'], r['score']])
        rn = ws2.max_row
        for ci in range(1, 5):
            ws2.cell(row=rn, column=ci).fill = fill
            ws2.cell(row=rn, column=ci).alignment = Alignment(
                horizontal='center' if ci not in (2, 3) else 'left')

# ── Sheet 3: 图例 ─────────────────────────────────────────────────────────────
ws3 = wb.create_sheet("图例")
ws3.column_dimensions['A'].width = 15
ws3.column_dimensions['B'].width = 30
ws3.column_dimensions['C'].width = 20

ws3.append(['颜色', '含义', '匹配得分范围'])
ws3.append(['绿色', '高置信度 — 匹配可靠', '≥ 0.70'])
ws3.append(['黄色', '中置信度 — 建议人工确认', '0.50 ~ 0.69'])
ws3.append(['红色', '低置信度 — 需人工确认', '< 0.50'])

ws3['A2'].fill = fill_high
ws3['A3'].fill = fill_medium
ws3['A4'].fill = fill_low

ws3.append([])
ws3.append(['统计', '', ''])
ws3.append([f'高置信度', f'{high} 个', f'{high/len(results)*100:.1f}%'])
ws3.append([f'中置信度', f'{medium} 个', f'{medium/len(results)*100:.1f}%'])
ws3.append([f'低置信度', f'{low} 个', f'{low/len(results)*100:.1f}%'])
ws3.append([f'总计', f'{len(results)} 个', '100%'])

wb.save(xlsx_path)
print(f"Excel 已保存: {xlsx_path}")
print("\n完成！")
